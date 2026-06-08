"""Generate downloadable DOCX files from markdown content."""
import io
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_markdown_paragraph(doc: Document, text: str) -> None:
    """Minimal markdown-to-docx conversion for headings and bullets."""
    if text.startswith("# "):
        p = doc.add_heading(text[2:], level=1)
    elif text.startswith("## "):
        p = doc.add_heading(text[3:], level=2)
    elif text.startswith("### "):
        p = doc.add_heading(text[4:], level=3)
    elif text.startswith("- ") or text.startswith("* "):
        doc.add_paragraph(text[2:], style="List Bullet")
    elif text.startswith("**") and text.endswith("**"):
        p = doc.add_paragraph()
        run = p.add_run(text[2:-2])
        run.bold = True
    else:
        doc.add_paragraph(text)


def markdown_to_docx(markdown_text: str) -> bytes:
    doc = Document()
    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for line in markdown_text.split("\n"):
        line = line.rstrip()
        if line:
            _add_markdown_paragraph(doc, line)
        else:
            doc.add_paragraph("")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
